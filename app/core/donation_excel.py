"""Lectura del Excel de insulinas que llega en una donación.

Función pura: recibe los bytes del .xlsx y devuelve las filas válidas y los
errores por fila. No toca la base de datos; la insulina se normaliza con la
función que pase quien la llama (el catálogo vive en `endpoints/donations.py`).

Columnas esperadas (sin importar mayúsculas ni acentos):
    CANTIDAD | PRODUCTO | TIPO DE INSULINA | PRESENTACIÓN | CONCENTRACIÓN
    y, opcionales, FECHA VENCIMIENTO | NRO LOTE.
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Callable, Dict, List, Optional

MAX_ROWS = 5000

_HEADER_ALIASES: Dict[str, tuple] = {
    "cantidad": ("cantidad", "cant", "envases"),
    "producto": ("producto", "nombre comercial", "nombre"),
    "tipo": ("tipo de insulina", "tipo insulina", "tipo", "insulina", "principio activo"),
    "presentacion": ("presentacion",),
    "concentracion": ("concentracion",),
    "vencimiento": ("fecha vencimiento", "fecha de vencimiento", "vencimiento", "fecha venc", "vence"),
    "lote": ("nro lote", "numero de lote", "nro de lote", "n lote", "lote"),
}
_REQUIRED = ("cantidad", "producto", "tipo", "presentacion", "concentracion")


class DonationExcelError(ValueError):
    """El archivo no se puede leer o no tiene las columnas necesarias."""


@dataclass
class DonationRow:
    row_number: int
    cantidad: int
    producto: str
    insulina: str  # nombre canónico del catálogo
    tipo_original: str
    presentacion_ml: float
    concentracion_ui_ml: int
    fecha_venc: Optional[date]
    lote: Optional[str]

    @property
    def ui_per_unit(self) -> float:
        return self.presentacion_ml * self.concentracion_ui_ml


@dataclass
class RowError:
    row_number: int
    producto: str
    message: str


@dataclass
class ParsedDonation:
    rows: List[DonationRow] = field(default_factory=list)
    errors: List[RowError] = field(default_factory=list)
    total_rows: int = 0


def _norm(value) -> str:
    text = unicodedata.normalize("NFD", str(value if value is not None else "").strip().lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text)).strip()


def _map_headers(header_cells) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for idx, cell in enumerate(header_cells):
        name = _norm(cell)
        if not name:
            continue
        for key, aliases in _HEADER_ALIASES.items():
            if key not in mapping and name in aliases:
                mapping[key] = idx
    missing = [k for k in _REQUIRED if k not in mapping]
    if missing:
        raise DonationExcelError(
            "Faltan columnas en el Excel: " + ", ".join(m.upper() for m in missing)
            + ". Se esperan CANTIDAD, PRODUCTO, TIPO DE INSULINA, PRESENTACIÓN y CONCENTRACIÓN."
        )
    return mapping


def _to_int(value) -> int:
    if isinstance(value, bool):
        raise ValueError
    if isinstance(value, (int, float)):
        if float(value) != int(value):
            raise ValueError
        return int(value)
    text = str(value).strip().replace(",", ".")
    number = float(text)
    if number != int(number):
        raise ValueError
    return int(number)


def _parse_ml(value) -> float:
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*ml", str(value or "").lower())
    if not match:
        if isinstance(value, (int, float)) and value > 0:
            return float(value)
        raise ValueError
    ml = float(match.group(1).replace(",", "."))
    if ml <= 0:
        raise ValueError
    return ml


def _parse_concentration(value) -> int:
    if isinstance(value, (int, float)) and value > 0:
        return int(value)
    match = re.search(r"(\d+)", str(value or ""))
    if not match or int(match.group(1)) <= 0:
        raise ValueError
    return int(match.group(1))


def _parse_date(value) -> Optional[date]:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError


def parse_donation_excel(
    content: bytes,
    normalize_insulin: Callable[[Optional[str]], Optional[str]],
) -> ParsedDonation:
    """Lee la primera hoja del .xlsx. Una fila con problemas no detiene a las demás."""
    try:
        from openpyxl import load_workbook  # import diferido: no es obligatorio para arrancar la API
    except ImportError as exc:  # pragma: no cover - depende del entorno
        raise DonationExcelError("Falta la dependencia 'openpyxl' en el servidor.") from exc

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise DonationExcelError("El archivo no es un Excel (.xlsx) válido.") from exc

    try:
        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration as exc:
            raise DonationExcelError("El Excel está vacío.") from exc
        columns = _map_headers(header)

        result = ParsedDonation()
        for row_number, cells in enumerate(rows_iter, start=2):
            if row_number - 1 > MAX_ROWS:
                raise DonationExcelError(f"El Excel supera el máximo de {MAX_ROWS} filas.")
            if not cells or all(c is None or str(c).strip() == "" for c in cells):
                continue
            result.total_rows += 1

            def cell(key):
                idx = columns.get(key)
                return cells[idx] if idx is not None and idx < len(cells) else None

            producto = str(cell("producto") or "").strip()
            try:
                try:
                    cantidad = _to_int(cell("cantidad"))
                except (ValueError, TypeError):
                    raise ValueError("CANTIDAD debe ser un número entero")
                if cantidad <= 0:
                    raise ValueError("CANTIDAD debe ser mayor a 0")
                if not producto:
                    raise ValueError("PRODUCTO está vacío")
                tipo_original = str(cell("tipo") or "").strip()
                insulina = normalize_insulin(tipo_original) or normalize_insulin(producto)
                if not insulina:
                    raise ValueError(f"«{tipo_original or producto}» no corresponde al catálogo de insulinas")
                try:
                    ml = _parse_ml(cell("presentacion"))
                except (ValueError, TypeError):
                    raise ValueError(f"PRESENTACIÓN «{cell('presentacion')}» no tiene un volumen en ml")
                try:
                    conc = _parse_concentration(cell("concentracion"))
                except (ValueError, TypeError):
                    raise ValueError(f"CONCENTRACIÓN «{cell('concentracion')}» no es válida (ej. U100)")
                try:
                    fecha = _parse_date(cell("vencimiento"))
                except (ValueError, TypeError):
                    raise ValueError(f"FECHA VENCIMIENTO «{cell('vencimiento')}» no es una fecha válida")
                lote_raw = cell("lote")
                lote = str(lote_raw).strip() if lote_raw is not None and str(lote_raw).strip() else None
            except ValueError as exc:
                result.errors.append(RowError(row_number, producto, str(exc)))
                continue

            result.rows.append(
                DonationRow(
                    row_number=row_number,
                    cantidad=cantidad,
                    producto=producto,
                    insulina=insulina,
                    tipo_original=tipo_original,
                    presentacion_ml=ml,
                    concentracion_ui_ml=conc,
                    fecha_venc=fecha,
                    lote=lote,
                )
            )
        return result
    finally:
        workbook.close()
